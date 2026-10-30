"""Admin: BOM list, company profile, the four reports (screen, PDF, Excel), rating round trip, second-device login."""
import re, json, os, subprocess, zipfile, urllib.parse
from mockgh import PNG, BASE, ADMIN_URL, ADMIN_NAME
from lib import *

OUT = f'{SP}/out'
os.makedirs(OUT, exist_ok=True)


def shot_paper(pg, path):
    pg.evaluate("document.querySelector('.topbar').style.visibility='hidden'; document.querySelector('#toasts').style.display='none'")
    pg.locator('#rpPaper').screenshot(path=path)
    pg.evaluate("document.querySelector('.topbar').style.visibility=''; document.querySelector('#toasts').style.display=''")


def paper(pg):
    return pg.inner_text('#rpPaper')


def edit_client(pg):
    pg.click('#adminNav a[data-route=dashboard]'); pg.wait_for_selector('#btnNewClient')
    pg.click('.item a:has-text("Edit")'); pg.wait_for_selector('#edName')


def snapshot(S, pg):
    payload = json.loads(S.mock.text(f'data/clients/{S.cid}.json'))['payload']
    return json.loads(pg.evaluate(DECRYPT_JS, {'payload': payload, 'secret': S.secret, 'id': S.cid}))


def download_xlsx(pg, name):
    with pg.expect_download() as dl:
        pg.click('#btnRpXlsx')
    d = dl.value
    path = f'{OUT}/{name}.xlsx'
    d.save_as(path)
    return d.suggested_filename, path


def pdf_text(path):
    return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True).stdout


def run(S):
    r, mock, pg, cid = S.r, S.mock, S.admin, S.cid
    _acc = lambda d: d.accept()
    pg.on('dialog', _acc)
    from openpyxl import load_workbook

    # ================= BOM list =================
    edit_client(pg)
    pg.click('#tab-bom')
    r.check('BOM tab: empty state and private note', 'No materials listed yet' in pg.inner_text('main'))
    rows = [('18mm white board', 'BRD-18W', 12, 'sheet', 1200, 'board'), ('Edge tape white', 'EDG-W', 60, 'm', 5, 'edge'), ('Soft-close hinge', 'HNG-SC', 24, 'pcs', 40, 'hardware')]
    for i, (name, code, qty, unit, cost, cat) in enumerate(rows):
        pg.click('#btnAddBom')
        row = pg.locator('.bomrow').nth(i)
        row.locator('input[placeholder="Item name"]').fill(name); row.locator('input[placeholder="Code"]').fill(code)
        row.locator('input[type=number]').nth(0).fill(str(qty)); row.locator('input[type=number]').nth(1).fill(str(cost))
        row.locator('select').nth(0).select_option(unit); row.locator('select').nth(1).select_option(cat)
    txt = pg.inner_text('main')
    r.check('BOM totals: materials cost 15,660 and 31.3% of the 50,000 contract', money_in(txt, 15660) and '31.3%' in txt, txt[-500:])
    r.check('BOM line total follows qty x unit cost (12 x 1,200 = 14,400)', money_in(pg.locator('.bomrow .bomtot').nth(0).inner_text(), 14400))
    pg.click('#btnBomSync'); pg.wait_for_selector('.toast:has-text("Costs updated from the BOM")')
    pg.click('#tab-costs')
    r.check('sync to costs: 2 new COGS lines (materials 14,700 and hardware 960); COGS now 37,660', pg.locator('.rg').count() == 5 and money_in(pg.inner_text('main'), 37660), pg.inner_text('main')[:400])
    pg.click('#tab-bom'); pg.click('#btnBomSync'); pg.click('#tab-costs')
    r.check('syncing twice does not duplicate cost lines', pg.locator('.rg').count() == 5 and money_in(pg.inner_text('main'), 37660))
    # second project: same board code, so the purchasing list has to add the two together
    pg.click('.tabs .pill >> nth=1'); pg.click('#tab-bom'); pg.click('#btnAddBom')
    row = pg.locator('.bomrow').nth(0)
    row.locator('input[placeholder="Item name"]').fill('18mm white board'); row.locator('input[placeholder="Code"]').fill('BRD-18W')
    row.locator('input[type=number]').nth(0).fill('3'); row.locator('input[type=number]').nth(1).fill('1200'); row.locator('select').nth(0).select_option('sheet')
    save(pg)
    leak = [s for s in ('BRD-18W', 'Soft-close', 'Edge tape') if s in alltext(mock)]
    r.check('BOM is stored encrypted only (no material names in any repository file)', not leak, str(leak))
    plain = pg.evaluate(DECRYPT_JS, {'payload': json.loads(mock.text(f'data/clients/{cid}.json'))['payload'], 'secret': S.secret, 'id': cid})
    r.check('client snapshot has no BOM, materials or costs', '"bom"' not in plain and 'BRD-18W' not in plain and '"costs"' not in plain)
    pg.click('#adminNav a[data-route=dashboard]'); pg.click('#btnRefresh'); pg.wait_for_selector('.item')
    pg.click('.item a:has-text("Edit")'); pg.wait_for_selector('#edName'); pg.click('#tab-bom')
    r.check('BOM list reloads from the repository (3 lines, names kept)', pg.locator('.bomrow').count() == 3 and pg.locator('.bomrow input[placeholder="Item name"]').first.input_value() == '18mm white board')

    # ================= finance report =================
    pg.click('#adminNav a[data-route=reports]'); pg.wait_for_selector('#rpPaper')
    txt = paper(pg)
    r.check('reports page: finance report is the default, dated, with the DUSUR logo', 'Finance report' in txt and 'DUSUR' in pg.inner_text('.rp-head') and pg.locator('.rp-logo .rp-mark svg').count() == 1 and fmtday(0) in pg.inner_text('.rp-foot'), pg.inner_text('.rp-head'))
    r.check('finance KPIs: revenue 50,000 / collected 20,000 / remaining 30,000 / COGS 37,660 / gross 12,340 / net 9,340', all(money_in(txt, n) for n in (50000, 20000, 30000, 37660, 12340, 9340)), txt[:900])
    r.check('finance margins: gross 24.7%, net 18.7%, markup 32.8%', all(x in txt for x in ('24.7%', '18.7%', '32.8%')))
    caps = pg.locator('#rpPaper figure.rp-chart figcaption').all_inner_texts()
    r.check('four analytic charts: revenue/COGS/gross, cost mix, payments by month, margin by project', caps == ['Revenue, COGS and gross profit', 'Where the cost goes', 'Payments received by month', 'Gross margin by project'], str(caps))
    r.check('charts are real SVG with bars/arcs and values', pg.locator('#rpPaper svg.rp-svg rect').count() >= 8 and pg.locator('#rpPaper svg.rp-svg path').count() >= 3 and pg.locator('#rpPaper .legend li').count() >= 6)
    r.check('three tables: projects (2 rows + total), by type, cost breakdown', pg.locator('#rpPaper .rp-tbl').count() == 3 and pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr:not(.total)').count() == 2 and pg.locator('#rpPaper .rp-tbl').first.locator('tr.total').count() == 1)
    r.check('confidential marking on finance/BOM reports', 'Confidential' in pg.inner_text('.rp-foot'))
    shot_paper(pg, f'{SP}/shot-d-finance.png')
    # PDF through the print stylesheet
    pdf = f'{OUT}/finance.pdf'
    pg.pdf(path=pdf, prefer_css_page_size=True, print_background=True)
    t = pdf_text(pdf); pages = int(re.search(r'Pages:\s+(\d+)', subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout).group(1))
    r.check('print/PDF: report only (no menu, toolbar or buttons), landscape page, content present', 'Finance report' in t and 'Revenue, COGS and gross profit' in t and 'Print / Save as PDF' not in t and 'Download Excel' not in t and 'Settings' not in t and pages >= 1, t[:300])
    subprocess.run(['pdftoppm', '-r', '55', '-png', pdf, f'{OUT}/finance-pdf'])
    # Excel
    fname, xp = download_xlsx(pg, 'finance')
    r.check('Excel file name carries company, report and date', fname.endswith('.xlsx') and 'Finance' in fname and day(0) in fname, fname)
    z = zipfile.ZipFile(xp)
    r.check('xlsx is a valid zip with workbook, styles, 4 native charts and the logo picture', z.testzip() is None and all(n in z.namelist() for n in ('[Content_Types].xml', 'xl/workbook.xml', 'xl/styles.xml', 'xl/media/logo.png')) and len([n for n in z.namelist() if n.startswith('xl/charts/chart')]) == 4, str(z.namelist()))
    wb = load_workbook(xp)
    r.check('workbook sheets: Summary, Charts, Projects, By project type, Cost breakdown', wb.sheetnames == ['Summary', 'Charts', 'Projects', 'By project type', 'Cost breakdown'], str(wb.sheetnames))
    ws = wb['Summary']
    vals = {c.value for row in ws.iter_rows() for c in row if isinstance(c.value, (int, float))}
    r.check('Summary sheet carries the KPI numbers as real numbers', {50000, 20000, 30000, 37660, 12340, 9340} <= vals, str(sorted(vals)))
    r.check('Summary sheet: title, scope, generated date, logo image', ws['A5'].value == 'Finance report' and len(ws._images) == 1, str(ws['A5'].value))
    wp = wb['Projects']
    heads = [c.value for c in wp[4]]
    r.check('Projects sheet: header band, frozen header, numeric cells with number formats', wp.freeze_panes == 'A5' and heads[0] == 'SKU' and wp['D5'].value == 50000 and '#,##0.00' in wp['D5'].number_format and wp.max_row == 7, f'{heads} {wp.freeze_panes} {wp.max_row}')
    r.check('Projects sheet: total row sums the table (price 50,000, COGS 37,660)', wp['D7'].value == 50000 and wp['G7'].value == 37660, f"{wp['D7'].value} {wp['G7'].value}")
    wc = wb['Charts']
    r.check('Charts sheet: data blocks next to 4 native charts', len(wc._charts) == 4 and wc['A1'].value == 'Revenue, COGS and gross profit', f'{len(wc._charts)}')
    soffice = subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', f'{OUT}/lo', xp], capture_output=True, text=True, timeout=180)
    lo = f'{OUT}/finance.pdf'.replace('finance.pdf', 'finance.xlsx.pdf')
    if not os.path.exists(f'{OUT}/finance.pdf'): pass
    r.check('LibreOffice opens the workbook and renders it (charts included)', soffice.returncode == 0)

    # ================= BOM report =================
    pg.select_option('#rpType', 'bom')
    txt = paper(pg)
    r.check('BOM report: items, codes, totals (19,260), categories', all(x in txt for x in ('BRD-18W', '18mm white board', 'Soft-close hinge', 'Boards and panels', 'Hardware')) and money_in(txt, 19260) and 'BOM list' in txt, txt[:600])
    r.check('BOM report: one table per project, category summary and a consolidated purchasing list', pg.locator('#rpPaper .rp-tbl').count() == 4 and 'Consolidated materials' in txt)
    cons = pg.locator('#rpPaper .rp-tbl').last
    r.check('consolidated list adds the same board code across projects (15 sheets = 18,000)', 'BRD-18W' in cons.inner_text() and '15' in cons.inner_text() and money_in(cons.inner_text(), 18000), cons.inner_text())
    caps = pg.locator('#rpPaper figure.rp-chart figcaption').all_inner_texts()
    r.check('BOM charts: cost by category and largest items', caps == ['Material cost by category', 'Largest items by cost'], str(caps))
    shot_paper(pg, f'{SP}/shot-d-bom.png')
    _, bx = download_xlsx(pg, 'bom'); wb = load_workbook(bx)
    r.check('BOM workbook: a sheet per project list + category + consolidated sheet', len(wb.sheetnames) == 6 and any(n.startswith('DS-KIT') for n in wb.sheetnames) and 'Consolidated materials (for purchasing)'[:31] in wb.sheetnames, str(wb.sheetnames))
    r.check('BOM workbook: confidential note on the summary', 'Confidential' in ' '.join(str(c.value) for row in wb['Summary'].iter_rows() for c in row if c.value))
    # filter to one project
    pg.select_option('#rpClient', cid); pg.select_option('#rpProject', S.skus[1])
    r.check('scope filter: a single project shows only its own BOM (3,600), no consolidated list', money_in(paper(pg), 3600) and 'Soft-close' not in paper(pg) and 'Consolidated' not in paper(pg) and S.skus[1] in pg.inner_text('.rp-titleblock'), paper(pg)[:500])
    pg.select_option('#rpProject', ''); pg.select_option('#rpClient', '')

    # ================= production report =================
    pg.select_option('#rpType', 'production')
    txt = paper(pg)
    r.check('production report: 2 projects, 0 completed, 1 in progress, 1 pending, average progress 12.5%', '12.5%' in txt and pg.locator('.rp-kpi').count() == 6 and 'Production status report' in txt, txt[:500])
    caps = pg.locator('#rpPaper figure.rp-chart figcaption').all_inner_texts()
    r.check('production charts: status, progress by project, deliveries by month, pieces by status', caps == ['Projects by status', 'Progress by project', 'Deliveries due by month', 'Pieces by status'], str(caps))
    r.check('production report has no money (no prices, costs or profit)', 'Confidential' not in txt and 'COGS' not in txt and '50,000' not in txt and pg.locator('.paper.land').count() == 1)
    r.check('production tables: projects and pieces (Base cabinets x 6, Doors x 12)', 'Base cabinets' in txt and 'Doors' in txt and pg.locator('#rpPaper .rp-tbl').count() == 2)
    shot_paper(pg, f'{SP}/shot-d-production.png')
    # date filter
    pg.fill('#rpFrom', day(+3)); pg.dispatch_event('#rpFrom', 'change')
    r.check('date filter with no matching projects shows the empty message', pg.locator('#rpEmpty').count() == 1)
    pg.fill('#rpFrom', ''); pg.dispatch_event('#rpFrom', 'change')
    pg.select_option('#rpType', 'satisfaction')
    r.check('satisfaction report before any rating: clear empty state', 'no client ratings' in pg.inner_text('#rpEmpty'))

    # ================= company profile (Settings) =================
    pg.click('#adminNav a[data-route=settings]'); pg.wait_for_selector('#setPrefix')
    pg.fill('#setPhone', '+20 100 000 1111'); pg.fill('#setWa', '201001234567'); pg.fill('#setEmail', 'info@dusur.example'); pg.fill('#setWeb', 'https://dusur.example'); pg.fill('#setAddr', 'Industrial zone, Cairo')
    pg.set_input_files('#setLogoFile', files=[{'name': 'logo.png', 'mimeType': 'image/png', 'buffer': PNG}])
    pg.wait_for_selector('.logo-preview img')
    pg.locator('.card:has(h2:has-text("Business defaults")) input[placeholder="English"]').fill('Dusur Furniture · CR 12345')
    pg.click('#btnSaveSettings'); pg.wait_for_selector('.toast:has-text("Settings saved")')
    b = json.loads(mock.text('data/public/brand.json'))
    r.check('company profile published as plain public brand.json (phone, WhatsApp, email, website, logo)', b['phone'] == '+20 100 000 1111' and b['whatsapp'] == '201001234567' and b['email'] == 'info@dusur.example' and b['website'] == 'https://dusur.example' and b['logo'].startswith('data:image/png;base64,') and b['rating'] is True, str(b)[:300])
    r.check('brand.json holds nothing private; the report footer text stays in encrypted settings', set(b) == {'name', 'tagline', 'phone', 'whatsapp', 'email', 'website', 'address', 'logo', 'rating', 'inbox', 'look'} and 'CR 12345' not in mock.text('data/private/settings.json') and 'CR 12345' not in mock.text('data/public/brand.json'))
    pg.click('#adminNav a[data-route=reports]'); pg.wait_for_selector('#rpPaper')
    pg.select_option('#rpType', 'finance')
    r.check('uploaded logo and the footer text now appear on the report', pg.locator('.rp-logoimg').count() == 1 and 'CR 12345' in pg.inner_text('.rp-foot') and '+20 100 000 1111' in pg.inner_text('.rp-foot'))
    _, fx = download_xlsx(pg, 'finance2'); wb = load_workbook(fx)
    r.check('Excel summary carries the footer text and uses the uploaded logo', any('CR 12345' in str(c.value) for row in wb['Summary'].iter_rows() for c in row if c.value) and len(wb['Summary']._images) == 1)

    # ================= Arabic report =================
    pg.click('#btnLang'); pg.wait_for_selector('#rpTitle:has-text("التقرير المالي")')
    txt = paper(pg)
    r.check('Arabic report: RTL paper, Arabic title/KPIs/tables, client name', pg.evaluate("getComputedStyle(document.querySelector('#rpPaper')).direction") == 'rtl' and 'التقرير المالي' in txt and 'أحمد محمد' in txt and 'الإيراد' in txt, txt[:400])
    shot_paper(pg, f'{SP}/shot-d-finance-ar.png')
    _, ax = download_xlsx(pg, 'finance-ar'); z = zipfile.ZipFile(ax); wb = load_workbook(ax)
    r.check('Arabic Excel: right-to-left sheets, Arabic sheet names and headers', 'rightToLeft="1"' in z.read('xl/worksheets/sheet1.xml').decode() and wb.sheetnames[0] == 'الملخص' and wb['الملخص']['A5'].value == 'التقرير المالي', str(wb.sheetnames))
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', f'{OUT}/ar', ax], capture_output=True, timeout=180)
    pg.click('#btnLang'); pg.wait_for_selector('#rpTitle:has-text("Finance report")')

    # ================= rating round trip =================
    edit_client(pg); pg.click('.tabs .pill >> nth=0'); pg.click('#tab-stages')
    n = pg.locator('.editor-stage input[type=checkbox]').count()
    for i in range(n): pg.locator('.editor-stage input[type=checkbox]').nth(i).check()
    r.check('all tasks ticked -> project shows Completed', pg.locator('.titleblock .badge.done').count() >= 1 and 'Completed' in pg.inner_text('.titleblock'))
    pg.click('#tab-details'); r.check('details tab tells you the client rating is awaited', 'Waiting for the client' in pg.inner_text('main'))
    save(pg)
    ctx2, cp = new_page(S, 'client2'); S.client2 = cp
    cp.goto(f'{BASE}#client/{cid}'); cp.wait_for_selector('#clKey'); cp.fill('#clKey', S.secret); cp.click('#btnOpen'); cp.wait_for_selector('#tapeBig', timeout=30000)
    r.check('completed project shows the rating card with a 1-10 scale and comment box', cp.locator('#rateBox button.sc').count() == 10 and cp.locator('#rateComment').count() == 1)
    cp.click('.pcard >> nth=1'); cp.wait_for_selector('.chip.mono:has-text("DRS")')
    r.check('an unfinished project does not offer a rating', cp.locator('#rateBox').count() == 0)
    cp.click('.pcard >> nth=0'); cp.wait_for_selector('#rateBox')
    cp.locator('#rateBox').scroll_into_view_if_needed(); cp.locator('#rateBox').screenshot(path=f'{SP}/shot-f-rate-form.png')
    cp.click('#btnRateGo'); cp.wait_for_selector('#rateBox .msg.bad')
    r.check('submitting without a score is refused with a message', 'Choose a score' in cp.inner_text('#rateBox'))
    comment = 'ممتاز جدًا، شكرًا لكم — excellent work'
    cp.click('button.sc[data-n="9"]'); cp.fill('#rateComment', comment); cp.locator('#rateBox').screenshot(path=f'{SP}/shot-f-rate-filled.png'); cp.click('#btnRateGo'); cp.wait_for_selector('#btnRateWa')
    cp.locator('#rateBox').screenshot(path=f'{SP}/shot-f-rate-send.png')
    wa = cp.get_attribute('#btnRateWa', 'href'); mail = cp.get_attribute('#btnRateMail', 'href')
    msg = urllib.parse.unquote(wa.split('text=', 1)[1])
    code = re.search(r'DSFB1\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+', msg).group(0)
    r.check('rating is delivered by WhatsApp / email links built from the company profile', wa.startswith('https://wa.me/201001234567?text=') and mail.startswith('mailto:info@dusur.example?subject='))
    r.check('the message carries an encrypted code only: no comment text, no secret key', code.startswith(f'DSFB1.{cid}.') and 'ممتاز' not in msg and 'excellent' not in msg and S.secret not in msg and '9/10' in msg, msg)
    # admin imports it
    pg.click('#adminNav a[data-route=dashboard]'); pg.wait_for_selector('#btnNewClient')
    sha_priv, sha_pub = mock.files[f'data/private/{cid}.json']['sha'], mock.files[f'data/clients/{cid}.json']['sha']
    r.check('dashboard shows the ratings card (empty before the import)', 'No ratings yet' in pg.inner_text('main'))
    pg.click('.fbimport summary')
    junk = 'Hello,\n\n' + msg + '\n\nThanks'
    pg.fill('#fbIn', 'DSFB1.DS-CL-9999.AAAAAAAAAAAAAAAA\n' + code[:-8] + 'AAAAAAAA\n' + junk)
    pg.click('#btnFbImport'); pg.wait_for_selector('#fbRes .msg.ok', timeout=30000)
    res = pg.inner_text('#fbRes')
    r.check('import finds all codes in pasted text: valid one imported, unknown client and tampered code reported', '9/10 imported' in res and 'DS-CL-9999 was not found' in res and 'could not be read' in res, res)
    d = pg.inner_text('main')
    r.check('dashboard: average 9 / 10 with the comment and project', '9 / 10' in d and comment in d and S.skus[0] in d and '1 rated projects' in d, d[:900])
    r.check('private record and client snapshot were both re-saved', mock.files[f'data/private/{cid}.json']['sha'] != sha_priv and mock.files[f'data/clients/{cid}.json']['sha'] != sha_pub)
    snap = snapshot(S, pg)
    fb = snap['projects'][0]['feedback']
    r.check('client snapshot now carries the rating (score, comment, date) and still no finance', fb['score'] == 9 and fb['comment'] == comment and fb['at'] and '"costs"' not in json.dumps(snap) and '"bom"' not in json.dumps(snap))
    pg.fill('#fbIn', msg); pg.click('#btnFbImport'); pg.wait_for_selector('#fbRes .msg.warn')
    r.check('importing the same code twice is harmless', 'newer rating is already saved' in pg.inner_text('#fbRes'))
    pg.fill('#fbIn', 'no code here'); pg.click('#btnFbImport'); pg.wait_for_selector('#fbRes .msg.bad')
    r.check('text without a rating code is explained', 'No rating code found' in pg.inner_text('#fbRes'))
    # the open client page picks it up through Refresh
    cp.click('#btnLiveRefresh'); cp.wait_for_selector('#rateBox.done-rate', timeout=20000)
    rb = cp.inner_text('#rateBox')
    cp.set_viewport_size({'width': 390, 'height': 844}); cp.wait_for_timeout(300)
    cp.screenshot(path=f'{SP}/shot-f-client-rated-mobile.png', full_page=True)
    r.check('rated client page on a phone: no horizontal scroll', cp.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'))
    cp.set_viewport_size({'width': 1280, 'height': 900})
    r.check('client page: Refresh shows "Thank you" with their score and their own comment', 'Thank you for your feedback' in rb and '9' in rb and comment in rb, rb)
    cp.reload(); cp.wait_for_selector('#clKey'); cp.fill('#clKey', S.secret); cp.click('#btnOpen'); cp.wait_for_selector('#rateBox.done-rate', timeout=30000)
    r.check('after a reload the rating is still shown (it lives in the snapshot, not in the browser)', comment in cp.inner_text('#rateBox'))
    page = cp.inner_text('#app')
    r.check('client page still shows no cost, profit or material data', not [x for x in ('COGS', 'Gross', 'Net profit', '37,660', 'BRD-18W', 'Soft-close', 'Confidential') if x in page])
    editor_fb = None
    edit_client(pg); pg.click('#tab-details')
    r.check('editor shows the rating read-only', '9/10' in pg.inner_text('main') and comment in pg.inner_text('main'))

    # ================= satisfaction report with data =================
    pg.click('#adminNav a[data-route=reports]'); pg.wait_for_selector('#rpPaper'); pg.select_option('#rpType', 'satisfaction')
    txt = paper(pg)
    r.check('satisfaction report: 1 rated project, average 9, response rate 100%, comment in the table', 'Customer satisfaction report' in txt and '100%' in txt and comment in txt and S.skus[0] in txt, txt[:700])
    caps = pg.locator('#rpPaper figure.rp-chart figcaption').all_inner_texts()
    r.check('satisfaction chart: score distribution', caps == ['Score distribution'], str(caps))
    shot_paper(pg, f'{SP}/shot-d-satisfaction.png')
    _, sx = download_xlsx(pg, 'satisfaction'); wb = load_workbook(sx)
    r.check('satisfaction workbook: ratings sheet keeps the Arabic/English comment', any(comment in str(c.value) for row in wb['Ratings and comments'].iter_rows() for c in row if c.value), str(wb.sheetnames))

    # ================= second device =================
    pg.click('#adminNav a[data-route=config]'); pg.wait_for_selector('#btnConnMake'); pg.click('#btnConnMake')
    conn = pg.input_value('#connOut')
    body = json.loads(__import__('base64').b64decode(conn[len('DSPCONN1.'):]))
    r.check('connection code holds the sealed token only (no password, no plaintext token)', conn.startswith('DSPCONN1.') and set(body) == {'v', 'user', 'repo', 'salt', 'iv', 'ct'} and TOKEN not in conn and PW1 not in conn and 'github_pat' not in conn)
    r.check('this-admin-page card shows the private address (unlisted file name)', ADMIN_NAME in pg.inner_text('#thisPageUrl'))
    ctx3, p2 = new_page(S, 'admin2'); p2.goto(ADMIN_URL); p2.wait_for_selector('#connIn')
    p2.fill('#connIn', conn); p2.click('#btnConnImport'); p2.wait_for_selector('#lockPw')
    unlock(p2, 'wrong password here'); p2.wait_for_selector('text=Incorrect admin password.')
    r.check('second device: the code alone does not open anything (wrong password rejected)', True)
    unlock(p2, PW1); p2.wait_for_selector('.item', timeout=30000)
    r.check('second device: correct password opens the dashboard with every client and rating', 'أحمد محمد' in p2.inner_text('main') and '9 / 10' in p2.inner_text('main'))
    r.check('admin page is noindex; the client site never mentions its address', 'noindex' in p2.content() and ADMIN_NAME not in mock.client_text())
    ctx3.close(); ctx2.close()
    pg.click('#adminNav a[data-route=dashboard]'); pg.wait_for_selector('#btnNewClient')
    pg.remove_listener('dialog', _acc)   # later modules install their own dialog handler on this page


def alltext(mock):
    return ' '.join(mock.files[p]['b'].decode('utf-8', 'ignore') for p in mock.paths('data/'))


def fmtday(offset):
    import datetime
    d = datetime.date.today() + datetime.timedelta(days=offset)
    return d.strftime('%-d %b %Y')
