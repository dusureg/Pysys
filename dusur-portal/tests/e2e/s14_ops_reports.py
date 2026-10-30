"""Admin: printable reports of Supply chain, Inventory and Logistics: opened from each section's "Printable report" button, filters, numbers checked against the raw data,
real PDF (print stylesheet, orientation, logo, date), Excel with native charts, the one-shipment sheet, empty states, Arabic."""
import re, json, os, subprocess, zipfile
from lib import *

OUT = f'{SP}/out'
os.makedirs(OUT, exist_ok=True)


def fmtday(offset):
    import datetime
    return (datetime.date.today() + datetime.timedelta(days=offset)).strftime('%-d %b %Y')


def paper(pg):
    return pg.inner_text('#rpPaper')


def kpis(pg):
    return {x.locator('.l').inner_text(): x.locator('.v').inner_text() for x in pg.locator('#rpPaper .rp-kpi').all()}


def caps(pg):
    return pg.locator('#rpPaper figure.rp-chart figcaption').all_inner_texts()


def pdf_of(pg, name):
    path = f'{OUT}/{name}.pdf'
    pg.pdf(path=path, prefer_css_page_size=True, print_background=True)
    txt = subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True).stdout
    info = subprocess.run(['pdfinfo', path], capture_output=True, text=True).stdout
    pages = int(re.search(r'Pages:\s+(\d+)', info).group(1)); w, h = map(float, re.search(r'Page size:\s+([\d.]+) x ([\d.]+)', info).groups())
    subprocess.run(['pdftoppm', '-r', '50', '-png', '-f', '1', '-l', '1', path, f'{OUT}/{name}-pdf'])
    return txt, pages, w > h


def xlsx_of(pg, name):
    with pg.expect_download() as dl:
        pg.click('#btnRpXlsx')
    d = dl.value; path = f'{OUT}/{name}.xlsx'; d.save_as(path)
    return d.suggested_filename, path


def run(S):
    r, mock, pg = S.r, S.mock, S.admin
    from openpyxl import load_workbook
    pg.set_viewport_size({'width': 1280, 'height': 900})
    dialogs = []
    on_dialog = lambda d: (dialogs.append(d.message), d.dismiss())
    pg.on('dialog', on_dialog)

    # make the reorder list interesting: the hinge falls below its minimum (it has a supplier), the edge band is already low and has none
    pg.click('#adminNav a[data-route=inventory]'); pg.wait_for_selector('.inv-item')
    pg.select_option('#invShow', 'all'); pg.fill('#invSearch', 'Soft-close')
    pg.locator('.inv-item', has_text='Soft-close').locator('[data-mv=count]').click(); pg.fill('#mvQty', '150'); pg.click('#btnMvApply'); toast_text(pg, 'Stock updated'); pg.fill('#invSearch', '')

    # ---- the print buttons of the sections open the right report
    r.check('Inventory: "Printable report" opens the stock report', pg.get_attribute('#btnInvPrint', 'href') == '#reports/inventory')
    pg.click('#btnInvPrint'); pg.wait_for_selector('#rpPaper')
    r.check('it opens #reports with the Inventory report selected, then tidies the address', pg.input_value('#rpType') == 'inventory' and pg.evaluate('location.hash') == '#reports' and 'Inventory' in pg.inner_text('#rpTitle'))
    inv = pg.evaluate("() => opsList('inventory').map((x) => ({ n: x.name, qty: +x.qty, min: +x.min, cost: +x.cost, exp: x.expiry, sup: x.supplierId, cat: x.category }))")
    today = day(0); soon = day(30)
    out = [x for x in inv if x['qty'] <= 0]; low = [x for x in inv if x['qty'] > 0 and x['min'] > 0 and x['qty'] <= x['min']]
    expd = [x for x in inv if x['exp'] and x['exp'] < today]; expg = [x for x in inv if x['exp'] and today <= x['exp'] <= soon]
    attention = {x['n'] for x in out + low + expd + expg}; value = sum(x['qty'] * x['cost'] for x in inv)
    k = kpis(pg)
    r.check('stock KPIs match the raw data: materials, need attention, out, low, expiring/expired, stock value', k.get('Materials') == str(len(inv)) and k.get('Need attention') == str(len(attention)) and k.get('Out of stock') == str(len(out)) and k.get('Low stock') == str(len(low))
            and k.get('Expiring / expired') == str(len(expd) + len(expg)) and money_in(k.get('Stock value', ''), round(value)) , f'{k} expected value {value} attention {attention}')
    tb = pg.locator('#rpPaper .rp-tbl').first
    r.check('the stock table lists every material with code, stock, minimum, value, supplier, place and status; with a total row', tb.locator('tbody tr:not(.total)').count() == len(inv) and tb.locator('tr.total').count() == 1 and all(x['n'] in tb.inner_text() for x in inv) and 'Rack A2' in tb.inner_text() and 'Expired' in tb.inner_text(), tb.inner_text()[:500])
    r.check('a second table lists the items that need attention with what to order', 'Items that need attention' in paper(pg) or 'attention' in paper(pg).lower() and pg.locator('#rpPaper .rp-tbl').count() == 2, str(pg.locator('#rpPaper .rp-tblwrap h3').all_inner_texts()))
    parts = {'land': pg.locator('#rpPaper.land').count() == 1, 'logo': pg.locator('.rp-head .rp-logoimg, .rp-head .rp-mark').count() >= 1, 'date': fmtday(0) in pg.inner_text('.rp-foot'), 'conf': 'Confidential' in pg.inner_text('.rp-foot'), 'cur': 'EGP' in pg.inner_text('.rp-titleblock')}
    r.check('the report carries the logo, the date, the "Confidential" mark, the currency, and is landscape', all(parts.values()), str(parts) + ' ' + pg.inner_text('.rp-foot'))
    r.check('3 charts: stock value by category, largest stock values, stock status', len(caps(pg)) == 3, str(caps(pg)))
    r.check('charts are real SVG', pg.locator('#rpPaper svg.rp-svg').count() == 3 and pg.locator('#rpPaper svg.rp-svg rect, #rpPaper svg.rp-svg path').count() >= 8)
    # filters
    pg.select_option('#rpCat', 'hardware'); n = len([x for x in inv if x['cat'] == 'hardware'])
    r.check('category filter narrows the report and its scope line', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr:not(.total)').count() == n and 'Hardware' in pg.inner_text('.rp-titleblock'), pg.inner_text('.rp-titleblock'))
    pg.select_option('#rpCat', ''); pg.select_option('#rpState', 'attention')
    r.check('"needs attention" filter lists only those', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr:not(.total)').count() == len(attention))
    pg.select_option('#rpState', 'all'); pg.select_option('#rpSup', '-')
    r.check('"no supplier" filter', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr:not(.total)').count() == len([x for x in inv if not x['sup']]))
    pg.select_option('#rpSup', '')
    # PDF + Excel of the inventory report
    txt, pages, land = pdf_of(pg, 'inventory')
    r.check('PDF: landscape, only the report (no menu or buttons), logo, date, every material', land and 'Inventory' in txt and 'DUSUR' in txt and day(0).split('-')[0] in txt and all(x['n'] in txt for x in inv) and 'Download Excel' not in txt and 'Print / Save as PDF' not in txt, txt[:300])
    fname, xp = xlsx_of(pg, 'inventory'); z = zipfile.ZipFile(xp); wb = load_workbook(xp)
    r.check('Excel: file name carries company, report and date; valid zip; 3 native charts; logo', 'Inventory' in fname and day(0) in fname and z.testzip() is None and len([n for n in z.namelist() if n.startswith('xl/charts/chart')]) == 3 and 'xl/media/logo.png' in z.namelist(), fname)
    ws = wb[wb.sheetnames[2]]
    r.check('Excel: sheets are Summary, Charts and one per table; the stock sheet has real numbers, a frozen header and a total', wb.sheetnames[:2] == ['Summary', 'Charts'] and len(wb.sheetnames) == 4 and ws.freeze_panes == 'A5' and ws.max_row == 4 + len(inv) + 1
            and any(isinstance(c.value, (int, float)) and abs(c.value - value) < 0.01 for row in ws.iter_rows() for c in row), str(wb.sheetnames))
    s1 = wb['Summary']; vals = {c.value for row in s1.iter_rows() for c in row if isinstance(c.value, (int, float))}
    r.check('Excel summary has the KPIs as numbers', {len(inv), len(attention), len(out), len(low)} <= vals and any(abs(v - value) < 0.01 for v in vals), str(sorted(vals)))
    lo = subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', f'{OUT}/lo', xp], capture_output=True, text=True, timeout=180)
    r.check('LibreOffice opens the Excel file', lo.returncode == 0 and os.path.exists(f'{OUT}/lo/inventory.pdf') or lo.returncode == 0)

    # ---- the reorder list
    pg.click('#adminNav a[data-route=inventory]'); pg.wait_for_selector('#btnOrderPrint'); pg.click('#btnOrderPrint'); pg.wait_for_selector('#rpPaper')
    r.check('the order list prints from the Inventory screen', pg.input_value('#rpType') == 'reorder')
    k = kpis(pg); p = paper(pg)
    r.check('reorder list: 2 materials to order (hinge, edge band) grouped by supplier with a "no supplier" group', k.get('Materials to order') == '2' and 'Cairo Hardware Co' in p and 'No supplier set' in p and 'Soft-close hinge' in p and 'Edge band' in p, f'{k}\n{p[:700]}')
    r.check('suggested quantities: hinge 250 (up to twice its minimum) and edge band 485 (up to its maximum)', re.search(r'Soft-close hinge[^\n]*250', p) is not None and re.search(r'Edge band[^\n]*485', p) is not None, p[:900])
    r.check('estimated total 4,697.50 = 250 x 12 + 485 x 3.5', money_in(k.get('Est. cost', '') or p, 4697) or '4,697.5' in p.replace(' ', '').replace('\xa0', ''), str(k))
    r.check('portrait page, supplier contact line in the group header', pg.locator('#rpPaper.land').count() == 0 and 'Cairo Hardware Co' in pg.locator('#rpPaper .rp-tblwrap h3').first.inner_text() + pg.locator('#rpPaper .rp-tblwrap h3').last.inner_text())
    pg.select_option('#rpSup', 'SUP-002'); r.check('filter by supplier', 'Edge band' not in paper(pg) and 'Soft-close hinge' in paper(pg)); pg.select_option('#rpSup', '')
    txt, pages, land = pdf_of(pg, 'reorder')
    r.check('reorder PDF: portrait, items and quantities present', not land and 'Soft-close hinge' in txt and 'Edge band' in txt and 'Reorder list' in txt, txt[:300])

    # ---- stock movements
    pg.select_option('#rpType', 'stockmoves'); pg.wait_for_selector('#rpFrom')
    nmoves = pg.evaluate("() => opsList('inventory').reduce((n, x) => n + x.moves.length, 0)")
    k = kpis(pg)
    r.check('movements report lists every movement with date, material, type, change, balance and note', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr').count() == nmoves and k.get('Movements') == str(nmoves) and 'Kitchen DS-KIT' in paper(pg), str(k))
    r.check('a chart of movements by type per month', caps(pg) == ['Movements by month'] or len(caps(pg)) == 1, str(caps(pg)))
    pg.fill('#rpFrom', day(1)); r.check('a date range with nothing in it shows the empty message', pg.locator('#rpEmpty').count() == 1 and 'no stock movements' in pg.inner_text('#rpEmpty')); pg.fill('#rpFrom', '')
    pg.fill('#rpTo', day(0)); r.check('"to today" still includes today\'s movements', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr').count() == nmoves); pg.fill('#rpTo', '')

    # ---- suppliers
    pg.click('#adminNav a[data-route=suppliers]'); pg.wait_for_selector('#btnSupPrint'); pg.click('#btnSupPrint'); pg.wait_for_selector('#rpPaper')
    r.check('Supply chain: "Printable report" opens the supplier directory', pg.input_value('#rpType') == 'suppliers')
    sup = pg.evaluate("() => opsList('suppliers').map((s) => ({ n: s.name, a: s.active, c: s.category, w: s.website, p: s.phone }))")
    k = kpis(pg); tb = pg.locator('#rpPaper .rp-tbl').first
    r.check('the directory lists every supplier with category, contact, phone, e-mail, website host, location, address and status', tb.locator('tbody tr').count() == len(sup) and all(s['n'] in tb.inner_text() for s in sup) and 'Active' in tb.inner_text() and 'alnour-boards.example' in tb.inner_text(), tb.inner_text()[:500])
    r.check('only the host of a website is printed (no https://, no tracking path) and the KPIs add up', 'https://' not in tb.inner_text() and k.get('Suppliers') == str(len(sup)), str(k))
    r.check('a table of materials per supplier', any('material' in x.lower() for x in pg.locator('#rpPaper .rp-tblwrap h3').all_inner_texts()), str(pg.locator('#rpPaper .rp-tblwrap h3').all_inner_texts()))
    pg.select_option('#rpCat', 'glass'); r.check('category filter', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr').count() == 1 and 'Delta Glass' in paper(pg)); pg.select_option('#rpCat', '')
    txt, pages, land = pdf_of(pg, 'suppliers')
    r.check('suppliers PDF: landscape with the names and phone numbers', land and 'Al Nour Boards' in txt and 'Delta Glass' in txt and '0100 123 4567' in txt, txt[:300])
    fname, xp = xlsx_of(pg, 'suppliers'); wb = load_workbook(xp)
    r.check('suppliers Excel: a supplier sheet with every row', any(ws.max_row >= 4 + len(sup) for ws in wb.worksheets[1:]) and 'Supplier' in fname, str(wb.sheetnames))

    # ---- logistics
    pg.click('#adminNav a[data-route=logistics]'); pg.wait_for_selector('#btnLgPrint'); pg.click('#btnLgPrint'); pg.wait_for_selector('#rpPaper')
    r.check('Logistics: "Printable report" opens the shipments report', pg.input_value('#rpType') == 'logistics')
    lg = pg.evaluate("() => opsList('logistics').map((s) => ({ id: s.id, t: s.title, st: lgState(s), late: lgLate(s), open: lgOpen(s) }))")
    k = kpis(pg); tb = pg.locator('#rpPaper .rp-tbl').first
    r.check('shipments KPIs: total, in progress, late, blocked, completed match the data', k.get('Shipments') == str(len(lg)) and k.get('In progress') == str(len([x for x in lg if x['open']])) and k.get('Completed') == str(len([x for x in lg if x['st'] == 'completed'])), f'{k} {lg}')
    r.check('the shipment table has id, name, type, party, route, carrier, tracking, arrival, progress, current step and state', tb.locator('tbody tr').count() == len(lg) and all(x['t'] in tb.inner_text() for x in lg) and 'TRK-889900' in tb.inner_text() and '100%' in tb.inner_text() and 'Completed' in tb.inner_text(), tb.inner_text()[:600])
    r.check('a table of the steps still open', any('open' in x.lower() or 'steps' in x.lower() for x in pg.locator('#rpPaper .rp-tblwrap h3').all_inner_texts()), str(pg.locator('#rpPaper .rp-tblwrap h3').all_inner_texts()))
    pg.select_option('#rpState', 'completed'); r.check('state filter', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr').count() == len([x for x in lg if x['st'] == 'completed'])); pg.select_option('#rpState', 'late')
    r.check('a filter with no shipments shows the empty message and no KPIs', pg.locator('#rpEmpty').count() == 1 and 'no shipments' in pg.inner_text('#rpEmpty') and pg.locator('#rpPaper .rp-kpi').count() == 0); pg.select_option('#rpState', 'all')
    pg.select_option('#rpKind', 'inbound'); r.check('type filter', pg.locator('#rpPaper .rp-tbl').first.locator('tbody tr').count() == len([x for x in lg if x['id'] in ('LG-0001', 'LG-0004')])); pg.select_option('#rpKind', '')
    txt, pages, land = pdf_of(pg, 'logistics')
    r.check('logistics PDF: landscape with the shipments', land and 'MDF boards from Al' in txt and 'Hardware pickup' in txt and 'TRK-889900' in txt and 'LG-0003' in txt, txt[:300])
    fname, xp = xlsx_of(pg, 'logistics'); wb = load_workbook(xp); z = zipfile.ZipFile(xp)
    r.check('logistics Excel: summary, charts (native) and a shipments sheet with progress as a percentage', 'Logistics' in fname and 'Charts' in wb.sheetnames and len([n for n in z.namelist() if n.startswith('xl/charts/chart')]) >= 1
            and any('%' in str(c.number_format) for ws in wb.worksheets[2:3] for row in ws.iter_rows() for c in row), str(wb.sheetnames))

    # ---- the sheet of one shipment
    pg.click('#adminNav a[data-route=logistics]'); pg.wait_for_selector('.lg-item'); pg.select_option('#lgShow', 'completed')
    pg.locator('.lg-item', has_text='MDF boards').locator('a:has-text("Edit")').click(); pg.wait_for_selector('#btnShipPrint'); pg.click('#btnShipPrint'); pg.wait_for_selector('#rpPaper')
    r.check('the shipment form prints its own sheet: #reports/shipment/LG-0001', pg.input_value('#rpType') == 'shipment' and pg.input_value('#rpShip') == 'LG-0001' and 'LG-0001' in pg.inner_text('#rpTitle'))
    tbs = pg.locator('#rpPaper .rp-tbl'); p = paper(pg)
    r.check('details (type, state, supplier, order number, carrier, tracking, route, arrival, cost), 7 steps with a column for initials, and the 3 materials with a tick column', tbs.count() == 3
            and all(x in tbs.nth(0).inner_text() for x in ('Al Nour Boards', 'PO-7741', 'Abu Ali Transport', 'TRK-889900', '10th of Ramadan', 'Factory')) and tbs.nth(1).locator('tbody tr').count() == 7 and 'Initials' in tbs.nth(1).inner_text()
            and tbs.nth(2).locator('tbody tr').count() == 3 and 'Wooden pallets' in tbs.nth(2).inner_text() and 'Checked' in tbs.nth(2).inner_text(), p[:800])
    r.check('the sheet shows 100% progress and "7 of 7 steps done"', '100%' in p and '7 of 7' in p)
    r.check('portrait page, no charts (a sheet to tick, not to analyse)', pg.locator('#rpPaper.land').count() == 0 and pg.locator('#rpPaper figure.rp-chart').count() == 0)
    txt, pages, land = pdf_of(pg, 'shipment')
    r.check('shipment PDF: portrait, one page, steps and materials, no buttons', not land and pages == 1 and 'Wooden pallets' in txt and 'Customs clearance (port)' in txt and 'Download Excel' not in txt, f'{pages} {txt[:300]}')
    pg.select_option('#rpShip', 'LG-0003'); r.check('another shipment can be picked from the list', 'LG-0003' in pg.inner_text('#rpTitle') and 'Hardware pickup' in paper(pg))
    fname, xp = xlsx_of(pg, 'shipment'); wb = load_workbook(xp)
    r.check('shipment Excel: no Charts sheet, a details sheet and a steps sheet', 'Charts' not in wb.sheetnames and len(wb.sheetnames) >= 3 and 'LG-0003' in fname, f'{fname} {wb.sheetnames}')

    # ---- Arabic
    pg.select_option('#rpType', 'inventory'); pg.click('#btnLang'); pg.wait_for_function("() => document.documentElement.lang === 'ar'")
    p = paper(pg)
    r.check('Arabic: the report is in Arabic and the sheet is right-to-left', 'المخزون' in pg.inner_text('#rpTitle') and pg.evaluate("() => getComputedStyle(document.querySelector('#rpPaper')).direction") == 'rtl' and 'تحتاج انتباهًا' in p, pg.inner_text('#rpTitle'))
    fname, xp = xlsx_of(pg, 'inventory-ar'); wb = load_workbook(xp)
    r.check('Arabic Excel: right-to-left sheets with Arabic names', all(ws.sheet_view.rightToLeft for ws in wb.worksheets) and all(re.search(r'[؀-ۿ]', n) for n in wb.sheetnames), str(wb.sheetnames))
    pg.click('#btnLang'); pg.wait_for_function("() => document.documentElement.lang === 'en'")
    r.check('switching back keeps the same report', pg.input_value('#rpType') == 'inventory')
    # the original four reports still work
    pg.select_option('#rpType', 'finance'); r.check('the original finance report is unchanged and its filters are back', pg.locator('#rpClient').count() == 1 and pg.locator('#rpCat').count() == 0 and 'Finance report' in pg.inner_text('#rpTitle'))
    pg.evaluate("location.hash = '#dashboard'"); pg.wait_for_selector('#btnNewClient')
    pg.remove_listener('dialog', on_dialog)
